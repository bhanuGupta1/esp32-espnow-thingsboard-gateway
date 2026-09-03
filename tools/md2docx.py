"""Convert a project Markdown document to .docx.

Not a general Markdown converter - it handles exactly the constructs these
documents use: headings, tables, fenced code, images, bullets, numbered lists,
blockquotes, rules, and inline bold/code.

    python tools/md2docx.py REPORT.md REPORT.docx
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

SRC = sys.argv[1]
OUT = sys.argv[2]
BASE = os.path.dirname(os.path.abspath(SRC))

# Portrait phone photos need a tighter cap than landscape screenshots, or a
# single figure eats a whole page.
PHOTO_WIDTH = Inches(3.6)
SCREENSHOT_WIDTH = Inches(6.0)


def add_runs(paragraph, text):
    """Emit inline **bold** and `code` as separate runs."""
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)


def main():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    with open(SRC, encoding="utf-8") as handle:
        lines = handle.read().split("\n")

    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Every rule in these documents separates a top-level section.
        if stripped == "---":
            doc.add_page_break()
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(6)
            run = para.add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            continue

        img = re.match(r"!\[(.*?)\]\((.+?)\)", stripped)
        if img:
            rel = img.group(2)
            path = os.path.join(BASE, rel.replace("/", os.sep))
            if os.path.exists(path):
                width = PHOTO_WIDTH if "photo_" in rel else SCREENSHOT_WIDTH
                doc.add_picture(path, width=width)
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph("[missing image: %s]" % rel)
            i += 1
            continue

        head = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if head:
            level = len(head.group(1))
            doc.add_heading(head.group(2).replace("**", ""), level=min(level, 4))
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < n and \
                re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for idx, cell_text in enumerate(header):
                cell = table.rows[0].cells[idx]
                cell.text = ""
                add_runs(cell.paragraphs[0], cell_text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for idx in range(min(len(row), len(header))):
                    cells[idx].text = ""
                    add_runs(cells[idx].paragraphs[0], row[idx])
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after = Pt(2)
                        for run in para.runs:
                            run.font.size = Pt(9.5)
            doc.add_paragraph()
            continue

        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.35)
            add_runs(para, " ".join(buf))
            for run in para.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x80, 0x30, 0x00)
            continue

        bullet = re.match(r"^([-*])\s+(.*)", stripped)
        number = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if bullet or number:
            text = (bullet or number).group(2)
            i += 1
            while i < n and lines[i].startswith("  ") and lines[i].strip() and \
                    not re.match(r"^\s*([-*]|\d+\.)\s", lines[i]):
                text += " " + lines[i].strip()
                i += 1
            para = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_runs(para, text)
            continue

        # Paragraph: join hard-wrapped lines until a blank or a new construct.
        buf = [stripped]
        i += 1
        while i < n and lines[i].strip() and not re.match(
            r"^(\||#{1,4}\s|```|>|!\[|\s*([-*]|\d+\.)\s|---$)", lines[i].strip()
        ):
            buf.append(lines[i].strip())
            i += 1
        text = " ".join(buf)
        para = doc.add_paragraph()
        if text.startswith("**Figure"):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(14)
            add_runs(para, text)
            for run in para.runs:
                run.font.size = Pt(9.5)
        else:
            add_runs(para, text)

    doc.save(OUT)
    print("wrote %s" % OUT)
    print("paragraphs=%d tables=%d images=%d" % (
        len(doc.paragraphs), len(doc.tables), len(doc.inline_shapes)))


main()
