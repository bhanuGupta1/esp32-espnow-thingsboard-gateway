"""Write speaker notes into each slide of the presentation.

Notes are what the presenter says, not a second copy of the slide. Short for
slides that carry themselves; longer for the two diagrams, the security slide
and the live hardware section, where the explanation is the content.
"""

import io
from pptx import Presentation

NOTES = {
1: """~15 s. Do not read the slide.

SAY: "Project one. Two ESP32 boards. One of them has no internet connection at
all, and its readings still reach the cloud - because the other one carries
them for it."

Then move on. The title slide is not where marks are.""",

2: """~40 s.

SAY: "Imagine a sensor somewhere with no Wi-Fi - a basement, a field, a corner
of a building. Connecting it properly is expensive. Every device needs its own
network password and its own setup, and joining Wi-Fi burns battery for several
seconds before a single reading is sent. So we flipped it. One device does the
connecting. The others talk to it cheaply over short-range radio."

IF ASKED why joining is expensive: scan every channel for beacons, authenticate
against the university RADIUS server, associate, then request an IP by DHCP -
all before one useful byte moves. On battery the radio is the largest draw in
the device.""",

3: """~30 s. Say the limitation FIRST, before anyone asks.

SAY: "Two ESP32 boards, each with a temperature and humidity sensor. Board 1
measures and sends by radio. Board 2 listens, adds its own reading, and
publishes both to the cloud."

THEN: "With only two boards this is one hop - one sender, one receiver. We
designed the message so a third board could pass it along, but nothing is ever
actually forwarded here. It is mesh-ready, not a mesh."

Volunteering your own limits reads as confidence. Being caught out on them does
not.

BACKUP: ESP32-D0WD-V3, dual core 240 MHz, 520 KB RAM, 4 MB flash, one 2.4 GHz
radio. DHT11: 0-50 C plus/minus 2, 20-90 % RH plus/minus 5, 1 Hz max sample
rate. Substituted for the BMP280 in the brief - documented in the report.""",

4: """~1 minute. This is your strongest slide. Point at the middle column first.

SAY: "Three stages - the sensor, the middleman, the cloud. The one in the middle
is on both networks at the same time: the radio network and the Wi-Fi. That is
the whole trick. It holds both sets of keys, so Board 1 reaches the internet
without ever joining it."

THEN WALK THE THREE ARROWS:
  Orange  - the radio between the boards. Dashed, because there is no wire.
  Green   - data going up to the cloud.
  Purple  - instructions coming back down.

THE LINE WORTH MEMORISING:
  "Green means we can watch it. Purple means we can control it."

WHY THREE COLUMNS: the diagram is drawn as columns so the single crossing point
is obvious. That crossing point is the entire argument of the project.

THE THEORY UNDERNEATH: networks are built in layers. ESP-NOW works at the link
layer - no IP address, no routing, no concept of "somewhere else". It can only
reach a radio it can physically hear. MQTT works at the application layer, on
top of TCP and IP. Because the two protocols sit at different layers, one device
can bridge them. A bridge between layers is what the word "gateway" means.

IF ASKED why two arrows for MQTT: because that is the difference between a
device you can watch and a device you can operate. One connection, two
directions.

FOOTNOTES ON THE SLIDE protect you: the access point chooses the channel, and
ttl and hop_count exist but nothing is ever forwarded.""",

5: """~40 s. The circuit diagram answers a different question from the block
diagram: block is what moves, circuit is which pin connects to which.

SAY: "Three wires per board - power, ground, and data. Identical on both boards
apart from which pin the data goes into: GPIO4 on one, GPIO5 on the other.
Notice there is no wire between the two boards. Separate USB supplies, separate
grounds. Nothing crosses between them except radio."

WHY THE PINS DIFFER: they were wired at different times and nothing required
them to match. That is the honest answer - do not invent a reason.

IF ASKED where the pull-up resistor is: "These are three-pin breakout modules -
the resistor is already on the module. A bare DHT11 would need about 10 k on the
data line."

IF ASKED about the dashed line: that is ESP-NOW, encrypted with a shared key
pair. It is drawn dashed precisely because it is not a wire.""",

6: """~45 s. This slide shows engineering judgement rather than coding.

SAY: "A radio can only listen on one channel at a time. Board 2 has to join the
campus Wi-Fi to reach the cloud, and the Wi-Fi network decides which channel
that is. We do not get to pick it - we find out at runtime and tell Board 1 to
match. That decides the whole setup order: you cannot configure Board 1 until
Board 2 tells you where it landed."

THEN THE HONEST PART: "It also means that if the Wi-Fi moves us to a different
channel, our link breaks silently. Wi-Fi and the cloud still look perfectly
healthy - only the radio link is dead. It happened three times during
development, so we check before every demonstration."

WHY IT IS SILENT: every indicator a casual observer checks still passes. The
gateway keeps publishing its own sensor. Only the node's data stops.

BACKUP: eduroam runs access points on channels 1, 6 and 11. We observed the
gateway on all three.""",

7: """~30 s.

SAY: "Every message is 26 bytes. It carries a version, who it is from and to, a
sequence number, and the two readings. Small on purpose - the radio protocol
allows 250 bytes, and we use 26."

WHY VERSION AND TYPE: so the format can change later without both boards
silently disagreeing about what the bytes mean.

BACKUP: a compile-time assertion locks the structure at exactly 26 bytes, so the
two sketches cannot drift apart without the build failing.""",

8: """~45 s. This is a genuinely subtle piece of design - take the time.

SAY: "Every message is numbered so we can spot repeats. But each board also
picks a random ID when it powers on, and that part is essential."

THEN THE WHY: "Without it, a board restarting looks exactly like an attacker
replaying old messages. The numbers start low again, we would reject every one
of them, and the board would never come back - and nothing would look broken."

THREE CASES the logic distinguishes:
  new packet     - accept
  duplicate      - reject
  gap (loss)     - accept, do not wait

A jump forward is a gap, not a duplicate. If we demanded consecutive numbers we
would deadlock on the first lost packet, waiting forever for something that is
never coming.

BACKUP: key is the triple (src_id, boot_id, seq). The comparison is wrap-aware -
unsigned subtraction then a signed compare - so it stays correct when a 32-bit
counter overflows.""",

9: """~90 seconds. Worth 25 marks. Slow down. Tell it as a story.

SAY: "At first, anyone with a matching radio nearby could send us a fake
reading. The message says who it is from, but that is just a claim - anyone can
write anything in that field."

"Our first fix was to check the sender's hardware address. Someone reviewed it
and pointed out the flaw straight away: an attacker can simply set that address
to whatever they like. We were checking a value the attacker controls. That is a
name badge, not a passport."

"So we encrypted the link instead. Both boards share a key. Now a message only
works if the sender actually has that key. That is the difference between
claiming who you are and proving it."

THEN CLOSE THE LOOP HONESTLY, before they ask:
"The cloud connection is still unencrypted - the access token crosses in plain
text. That is the most serious weakness left, and the fix is TLS on port 8883."

THE RULE UNDERNEATH IT: you cannot authenticate using a value the attacker
supplies. That is the whole lesson and it generalises far beyond this project.

WHAT THE ATTACK ACTUALLY DID: not a wrong temperature on a dashboard. A forged
packet carrying a very high sequence number advances the duplicate-detection
baseline, and after that every genuine packet from the real sensor is rejected
as a duplicate until the board reboots. One forged frame silences the real
sensor. It is a denial of service delivered through the integrity mechanism
itself.

BACKUP: a primary master key protects the per-peer key during setup; the
per-peer key encrypts the payload. The receiver also needs a peer entry, or the
driver discards the frame before any application code runs - a silent failure
with no error at all.""",

10: """~30 s.

SAY: "634 messages received and validated. 82 repeats caught and rejected. Zero
corrupted - both boards agree on the format. Fourteen out of fourteen acceptance
criteria evidenced."

ON THE OFFLINE BEHAVIOUR: "When the node goes quiet we keep publishing its last
known temperature, so the chart stays continuous - but a separate flag says the
data is stale. We do not put 'unknown' in the data field. We put it in a status
field."

That is a design decision worth defending if asked: nobody is misled, and the
chart stays readable.""",

11: """~40 s.

SAY: "Two physically separate sensors appear in the cloud as one device and one
data stream, because the merge happens on Board 2 before anything leaves the
building. Blue is the gateway's own sensor. Green is the remote node, arriving
by radio."

WHY MERGE AT THE EDGE: less traffic, less power, and the cloud does not have to
correlate two separate streams. Two 26-byte frames every five seconds become one
publish every ten. And the remote node needs no cloud credentials at all.

IF THE TRACE LOOKS STEPPED: that is the DHT11's one-degree resolution, not a
frozen sensor. Say it before they wonder.""",

12: """~40 s.

SAY: "Telemetry alone makes a device observable. It does not make it
manageable. These three buttons send commands from the cloud down to the
device - change how often it reports, reset its counters, ask for its status."

THE KEY POINT: "And the device reports the change back in its own telemetry. So
we are not just seeing that the command was sent - we are seeing that it
arrived and took effect. The round trip is confirmed at both ends."

ON VALIDATION: out-of-range values are refused rather than quietly adjusted.
Silently applying something different from what was asked would leave the
operator believing something untrue about the device.""",

13: """3 MINUTES. Most of the marks are here. Stop presenting, start operating,
and keep talking while you do it.

BEFORE YOU BEGIN: three windows visible at once - Board 1 serial, Board 2
serial, ThingsBoard. The argument only lands if all three are on screen.

1. BOTH TERMINALS
   "Left is Board 1 sending. Right is Board 2 receiving. Same sequence number,
   one millisecond apart."
   If pressed on why that matters: the acknowledgement is from the receiving
   radio, not the application. It proves the bytes arrived, not that they were
   accepted. The arrival line on the other board proves acceptance. Together
   they prove both - which one port alone cannot.

2. THUMB ON THE SENSOR
   "Watch the temperature climb."
   WAIT. Let them see it move. This proves the entire chain - sensor, radio,
   gateway, internet, dashboard - in about fifteen seconds. It is the single
   most convincing thing you can do.

3. PRESS "FAST PUBLISH (2 s)"
   "That button sends an instruction from the cloud down to the device. It
   changes how often it reports - from ten seconds to two."
   Then show the value returning: "And there is the device confirming it. Not
   just that we sent it - that it arrived and took effect."

4. UNPLUG BOARD 1 - keep talking, do not stand in silence
   "The gateway keeps publishing its own reading. Give it twenty seconds and it
   will mark the other board offline - but it keeps showing the last known
   temperature so the chart does not gap. The status flag carries the truth
   about freshness."

5. PLUG IT BACK IN
   "New random ID, numbering restarts, accepted cleanly. No flood of rejected
   duplicates."

IF THE LINK IS DEAD: it is the channel. Read it off Board 2's boot banner, put
it in sketch 04, re-upload to Board 1. Two minutes. Do not guess.""",

14: """~30 s.

SAY: "The brief specified a BMP280. The hardware available was a DHT11. That
changes what we measure - humidity instead of pressure - but not the
architecture we are demonstrating. It is documented in the report rather than
glossed over."

Owning a substitution is stronger than hoping nobody checks.""",

15: """~45 s. Name your own limits. It reads as maturity, not weakness.

SAY: "One hop only - two radios cannot demonstrate forwarding. The cloud
connection is not encrypted, so the access token crosses in clear text.
Credentials are compiled into the firmware and recoverable over USB. And the
channel is set by hand, so a Wi-Fi roam breaks the link silently."

DO NOT list ESP-NOW encryption as a limitation. It is implemented - a primary
master key plus a per-peer key, and the receiver authenticates the sender. That
is a feature you built, not a gap.""",

16: """~30 s.

SAY: "We tested the assumptions rather than trusting them. Port 1883 is open
outbound on eduroam - we verified that rather than assuming it. The realm for
authentication turned out to be the university domain, not the student mail
domain, and enterprise Wi-Fi gives you no error that distinguishes a wrong
username from a wrong password or being out of range. We found it by changing
one thing at a time."

That story answers "what was the hardest part" before it is asked.""",

17: """~30 s.

SAY: "Next steps, in order of value: a third board so the forwarding fields are
actually exercised. TLS on the cloud leg, which is the most serious remaining
exposure. Credentials provisioned at first boot rather than compiled in. And the
node scanning for the gateway instead of being told the channel by hand - that
would have prevented three outages during development."

Each item maps to a limitation on the previous slide. That symmetry is
deliberate - show it.""",

18: """Questions - 5 minutes.

"HOW DOES IT WORK?" - One board measures and radios it across. The other adds
its own reading and sends both to the cloud.

"WHY NOT WI-FI ON BOTH?" - Both would need credentials, setup and battery to
join. One gateway carrying several nodes scales better.

"IS IT SECURE?" - The radio link is: encrypted, keys on both ends, and the
receiver rejects anything that does not decrypt. The cloud link is not. That is
the honest gap.

"IS IT A MESH?" - Mesh-ready. Two radios can only do one hop. The format carries
a hop counter; we simply never forward.

"WHY QoS 0?" - Telemetry every ten seconds. A lost sample is superseded ten
seconds later. QoS 1 costs round trips and a queue for data with a very short
shelf life.

"ISN'T THE ACKNOWLEDGEMENT PROOF ENOUGH?" - No. It comes from the radio, not the
application. It proves the bytes arrived, not that they were accepted. That is
why we captured both boards against one clock.

"WHAT WAS THE HARDEST PART?" - Not the code, the campus network. Enterprise
Wi-Fi gives no useful error, so a wrong username looks identical to a wrong
password or being out of range.

"WHY DHT11?" - It is what was available. It changes what we measure, not the
architecture.

IF YOU DO NOT KNOW: "I don't know off the top of my head - it is in the report
and I can check." That costs almost nothing. Guessing and being wrong costs a
great deal.""",
}

prs = Presentation("PRESENTATION.pptx")
written = 0
for i, slide in enumerate(prs.slides, start=1):
    text = NOTES.get(i)
    if not text:
        continue
    slide.notes_slide.notes_text_frame.text = text.strip()
    written += 1

prs.save("PRESENTATION.pptx")
print("notes written to %d of %d slides" % (written, len(prs.slides)))
